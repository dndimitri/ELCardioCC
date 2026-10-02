import argparse
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


DEFAULT_GOLD = "test_gold.jsonl"
DEFAULT_PARTICIPANTS_DIR = "participants"
DEFAULT_OUTPUT_XLSX = "evaluation_results.xlsx"


def load_jsonl(path):
    rows = []
    with open(path, "r", encoding="utf-8") as f:
        for line_no, line in enumerate(f, start=1):
            line = line.strip()
            if not line:
                continue
            try:
                rows.append(json.loads(line))
            except json.JSONDecodeError as e:
                raise ValueError(f"Invalid JSON in {path}, line {line_no}: {e}") from e
    return rows


def get_patient_id(sample, index):
    return str(sample.get("patient_id", index))


def normalize_code(code):
    return str(code).strip()


def validate_gold_annotations(annotations, patient_id):
    if not isinstance(annotations, list):
        raise ValueError(f"Gold annotations for patient {patient_id} must be a list of lists.")

    normalized = []
    for i, group in enumerate(annotations):
        if not isinstance(group, list):
            raise ValueError(
                f"Gold annotation group {i} for patient {patient_id} must be a list."
            )

        clean_group = [normalize_code(code) for code in group if normalize_code(code)]

        if not clean_group:
            print(f"Warning: empty gold group {i} for patient {patient_id} skipped")
            continue

        normalized.append(clean_group)

    return normalized


def validate_submission_annotations(annotations, patient_id):
    if not isinstance(annotations, list):
        raise ValueError(f"Submitted annotations for patient {patient_id} must be a flat list.")

    normalized = []
    for i, code in enumerate(annotations):
        if isinstance(code, list):
            raise ValueError(
                f"Submitted annotation {i} for patient {patient_id} is a list. "
                "Submission annotations must be flat lists."
            )

        clean_code = normalize_code(code)
        if clean_code:
            normalized.append(clean_code)

    return normalized


def compute_group_metrics_single(predictions, ground_truth):
    """
    Compute TP, FP, FN for a single document.

    Args:
        predictions: list of predicted ICD-10 code strings
        ground_truth: list of label groups, each a list of ICD-10 code strings
                      e.g. [["R55"], ["I10", "I11"], ["I44"], ["Z95"]]

    Returns:
        (tp, fp, fn) tuple
    """
    # All codes appearing in any ground truth group
    all_gt_codes = set()
    for group in ground_truth:
        for code in group:
            all_gt_codes.add(code)

    pred_set = set(predictions)

    # TP: number of GT groups matched (at least one code predicted)
    tp = 0
    for group in ground_truth:
        if any(code in pred_set for code in group):
            tp += 1

    # FN: unmatched GT groups
    fn = len(ground_truth) - tp

    # FP: predicted codes not in any GT group
    fp = sum(1 for p in pred_set if p not in all_gt_codes)

    return tp, fp, fn


def compute_prf(tp, fp, fn):
    """Compute precision, recall, F1 from aggregate counts."""
    precision = tp / (tp + fp) if (tp + fp) else 0.0
    recall = tp / (tp + fn) if (tp + fn) else 0.0
    f1 = (2 * precision * recall / (precision + recall) if (precision + recall) else 0.0)
    return round(precision, 4), round(recall, 4), round(f1, 4)


def align_by_patient_id(gold_rows, sub_rows):
    gold_ids = [get_patient_id(row, i) for i, row in enumerate(gold_rows)]
    sub_ids = [get_patient_id(row, i) for i, row in enumerate(sub_rows)]

    if len(gold_rows) == len(sub_rows) and gold_ids == sub_ids:
        return [
            (gold_ids[i], gold_rows[i], sub_rows[i])
            for i in range(len(gold_rows))
        ]

    if len(set(gold_ids)) != len(gold_ids):
        raise ValueError("Gold file contains duplicate patient_id values.")
    if len(set(sub_ids)) != len(sub_ids):
        raise ValueError("Submission file contains duplicate patient_id values.")

    gold_by_id = dict(zip(gold_ids, gold_rows))
    sub_by_id = dict(zip(sub_ids, sub_rows))

    missing = sorted(set(gold_by_id) - set(sub_by_id))
    extra = sorted(set(sub_by_id) - set(gold_by_id))

    if missing:
        raise ValueError(f"Submission is missing {len(missing)} patient(s), e.g. {missing[:5]}")
    if extra:
        raise ValueError(f"Submission contains {len(extra)} extra patient(s), e.g. {extra[:5]}")

    return [(pid, gold_by_id[pid], sub_by_id[pid]) for pid in gold_ids]


def evaluate_submission(gold_rows, submission_path):
    sub_rows = load_jsonl(submission_path)

    if not sub_rows:
        raise ValueError("Submission file is empty.")

    aligned = align_by_patient_id(gold_rows, sub_rows)

    total_tp = 0
    total_fp = 0
    total_fn = 0

    for patient_id, gold_row, sub_row in aligned:
        if "annotations" not in gold_row:
            raise KeyError(f"Gold row for patient {patient_id} has no 'annotations' key.")
        if "annotations" not in sub_row:
            raise KeyError(f"Submission row for patient {patient_id} has no 'annotations' key.")

        gt = validate_gold_annotations(gold_row["annotations"], patient_id)
        preds = validate_submission_annotations(sub_row["annotations"], patient_id)

        tp, fp, fn = compute_group_metrics_single(preds, gt)

        total_tp += tp
        total_fp += fp
        total_fn += fn

    precision, recall, f1 = compute_prf(total_tp, total_fp, total_fn)

    return {
        "precision": precision,
        "recall": recall,
        "f1": f1,
    }


def find_submissions(participants_dir):
    participants_dir = Path(participants_dir)

    if not participants_dir.exists():
        raise FileNotFoundError(f"Participants folder not found: {participants_dir}")

    submissions = []

    for team_dir in sorted(participants_dir.iterdir()):
        if not team_dir.is_dir():
            continue

        team_name = team_dir.name
        jsonl_files = sorted(team_dir.glob("*.jsonl"))

        for submission_path in jsonl_files:
            submissions.append({
                "team": team_name,
                "submission": submission_path.name,
                "path": submission_path,
            })

    return submissions


def write_excel(results, output_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"

    headers = ["Team", "Submission", "Recall", "Precision", "F1"]
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = Border(bottom=thin)

    successful = [r for r in results if r["status"] == "OK"]
    failed = [r for r in results if r["status"] != "OK"]

    successful.sort(
        key=lambda x: (
            -x["f1"],
            -x["recall"],
            -x["precision"],
            x["team"].lower(),
            x["submission"].lower(),
        )
)

    failed.sort(
        key=lambda x: (
            x["team"].lower(),
            x["submission"].lower(),
        )
    )

    for r in successful:
        ws.append([
            r["team"],
            r["submission"],
            r["recall"],
            r["precision"],
            r["f1"],
        ])

    if failed:
        ws.append([""] * len(headers))
        ws.append(["FAILED SUBMISSIONS", "", "", "", ""])

        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="FCE4D6")

        for r in failed:
            ws.append([
                r["team"],
                r["submission"],
                "",
                "",
                "",
            ])

    for row in ws.iter_rows(min_row=2, min_col=3, max_col=5):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.0000"

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    widths = {
        "A": 24,
        "B": 42,
        "C": 14,
        "D": 14,
        "E": 14,
    }

    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    wb.save(output_path)


def run_batch_evaluation(gold_path, participants_dir, output_xlsx):
    gold_rows = load_jsonl(gold_path)

    if not gold_rows:
        raise ValueError("Gold file is empty.")

    submissions = find_submissions(participants_dir)

    if not submissions:
        raise ValueError(f"No .jsonl submissions found under {participants_dir}")

    results = []

    print(f"Gold file: {gold_path}")
    print(f"Participants folder: {participants_dir}")
    print(f"Found {len(submissions)} submission(s).\n")

    for item in submissions:
        team = item["team"]
        submission_name = item["submission"]
        submission_path = item["path"]

        print(f"Evaluating {team} / {submission_name}...")

        try:
            metrics = evaluate_submission(gold_rows, submission_path)
            result = {
                "team": team,
                "submission": submission_name,
                "path": str(submission_path),
                "status": "OK",
                "error": "",
                **metrics,
            }
            print(
                f"  OK | "
                f"Recall={metrics['recall']:.6f} "
                f"Precision={metrics['precision']:.6f} "
                f"F1={metrics['f1']:.6f}"
            )

        except Exception as e:
            result = {
                "team": team,
                "submission": submission_name,
                "path": str(submission_path),
                "status": "ERROR",
                "error": str(e),
            }
            print(f"  ERROR | {e}")

        results.append(result)

    write_excel(results, output_xlsx)

    print()
    print("=" * 60)
    print(f"Saved Excel results to: {output_xlsx}")
    print("=" * 60)


def run_tests():
    tests = [
        (
            "perfect",
            ["R55", "I10", "I44", "Z95"],
            [["R55"], ["I10", "I11"], ["I44"], ["Z95"]],
            (4, 0, 0),
        ),
        (
            "extra alias from same group is not FP",
            ["R55", "I10", "I11", "I44", "Z95"],
            [["R55"], ["I10", "I11"], ["I44"], ["Z95"]],
            (4, 0, 0),
        ),
        (
            "missing one group",
            ["R55", "I44", "Z95"],
            [["R55"], ["I10", "I11"], ["I44"], ["Z95"]],
            (3, 0, 1),
        ),
        (
            "extra code outside GT union",
            ["R55", "I10", "I44", "Z95", "I20"],
            [["R55"], ["I10", "I11"], ["I44"], ["Z95"]],
            (4, 1, 0),
        ),
        (
            "overlap: one code covers two groups",
            ["I20"],
            [["I20"], ["I20", "R07"]],
            (2, 0, 0),
        ),
        (
            "overlap: R07 covers only second group",
            ["R07"],
            [["I20"], ["I20", "R07"]],
            (1, 0, 1),
        ),
        (
            "empty prediction",
            [],
            [["R55"], ["I10", "I11"]],
            (0, 0, 2),
        ),
    ]

    passed = 0

    for name, preds, gt, expected in tests:
        actual = compute_group_metrics_single(preds, gt)
        ok = actual == expected
        print(f"[{'PASS' if ok else 'FAIL'}] {name}: got={actual}, expected={expected}")
        passed += int(ok)

    print(f"\nTests passed: {passed}/{len(tests)}")
    return passed == len(tests)


def main():
    parser = argparse.ArgumentParser(
        description="Batch evaluate ICD-10 submissions and produce Excel results."
    )

    parser.add_argument(
        "--gold",
        default=DEFAULT_GOLD,
        help=f"Path to gold JSONL file. Default: {DEFAULT_GOLD}",
    )

    parser.add_argument(
        "--participants-dir",
        default=DEFAULT_PARTICIPANTS_DIR,
        help=f"Folder containing one folder per participant. Default: {DEFAULT_PARTICIPANTS_DIR}",
    )

    parser.add_argument(
        "--output",
        default=DEFAULT_OUTPUT_XLSX,
        help=f"Output Excel file. Default: {DEFAULT_OUTPUT_XLSX}",
    )

    parser.add_argument(
        "--test",
        action="store_true",
        help="Run unit tests before evaluation.",
    )

    args = parser.parse_args()

    if args.test:
        success = run_tests()
        if not success:
            sys.exit(1)
        print()

    run_batch_evaluation(
        gold_path=args.gold,
        participants_dir=args.participants_dir,
        output_xlsx=args.output,
    )


if __name__ == "__main__":
    main()